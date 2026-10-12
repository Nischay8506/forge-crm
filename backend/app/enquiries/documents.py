"""Reading what is attached to an enquiry (decision N10 A).

PDFs with text and Excel sheets are read directly; Word files and plain text too.
Scanned PDFs, photos and drawings carry no text, so their images go to the vision
model (Gemini) when the AI key is set. CAD files and anything else are kept, not read.

Attachments come from strangers, so each reader has limits (pages, rows, time) and
any failure only means "not read", never a failed email.
"""

import asyncio
import csv
import io
import logging
import re
import zipfile
from dataclasses import dataclass, field

log = logging.getLogger(__name__)

MAX_PDF_PAGES = 30
MAX_ROWS = 500
MAX_COLS = 30
MAX_CHARS = 20_000
MAX_IMAGES = 5
MAX_IMAGE_BYTES = 4 * 1024 * 1024
READ_SECONDS = 20
IMAGE_TYPES = {"image/png", "image/jpeg", "image/jpg", "image/webp", "image/gif"}
CAD = (".dwg", ".dxf", ".step", ".stp", ".igs", ".iges", ".sldprt", ".prt")


@dataclass
class DocText:
    filename: str
    kind: str  # pdf | excel | csv | word | text | scan | image | cad | other
    text: str | None = None
    images: list[tuple[str, bytes]] = field(default_factory=list)  # (mime, bytes)
    note: str | None = None  # why it was not (fully) read
    file_id: str | None = None

    @property
    def needs_vision(self) -> bool:
        return bool(self.images) and not self.text


def _ext(name: str) -> str:
    m = re.search(r"\.[A-Za-z0-9]{1,6}$", name or "")
    return m.group(0).lower() if m else ""


def _clip(text: str) -> str:
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    return text[:MAX_CHARS]


def _pdf(data: bytes) -> tuple[str, list[tuple[str, bytes]], str | None]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception:  # noqa: BLE001
            return "", [], "The PDF is password-protected"
    pages = reader.pages[:MAX_PDF_PAGES]
    text = "\n".join((p.extract_text() or "") for p in pages)
    note = (f"Only the first {MAX_PDF_PAGES} pages were read"
            if len(reader.pages) > MAX_PDF_PAGES else None)
    images: list[tuple[str, bytes]] = []
    if len(text.strip()) < 40:  # a scan: the pages are pictures
        for p in pages:
            for img in p.images:
                mime = "image/png" if img.name.lower().endswith(".png") else "image/jpeg"
                if len(img.data) <= MAX_IMAGE_BYTES:
                    images.append((mime, img.data))
                if len(images) >= MAX_IMAGES:
                    break
            if len(images) >= MAX_IMAGES:
                break
    return text, images, note


def _excel(data: bytes) -> tuple[str, str | None]:
    import defusedxml  # noqa: F401 - openpyxl uses it to read the sheets' XML safely
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out, cut = [], False
    for ws in wb.worksheets[:10]:
        out.append(f"[Sheet: {ws.title}]")
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= MAX_ROWS:
                cut = True
                break
            cells = ["" if v is None else str(v).strip() for v in row[:MAX_COLS]]
            if any(cells):
                out.append(" | ".join(cells).rstrip(" |"))
    wb.close()
    return "\n".join(out), (f"Only the first {MAX_ROWS} rows of each sheet were read"
                            if cut else None)


def _csv(data: bytes) -> str:
    text = _decode(data)
    rows = []
    for i, row in enumerate(csv.reader(io.StringIO(text))):
        if i >= MAX_ROWS:
            break
        rows.append(" | ".join(c.strip() for c in row[:MAX_COLS]))
    return "\n".join(rows)


def _word(data: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        info = z.getinfo("word/document.xml")
        if info.file_size > 20 * 1024 * 1024:
            raise ValueError("too large")
        xml = z.read(info).decode("utf-8", "replace")
    xml = re.sub(r"</w:p>|<w:br/>|</w:tr>", "\n", xml)
    xml = re.sub(r"</w:tc>", " | ", xml)
    return re.sub(r"<[^>]+>", "", xml)


def _decode(data: bytes) -> str:
    for enc in ("utf-8", "utf-16", "cp1252"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", "replace")


def read_sync(filename: str, content_type: str, data: bytes) -> DocText:
    ext, ct = _ext(filename), (content_type or "").lower()
    try:
        if ext == ".pdf" or ct == "application/pdf":
            text, images, note = _pdf(data)
            if text.strip() and len(text.strip()) >= 40:
                return DocText(filename, "pdf", _clip(text), note=note)
            if images:
                return DocText(filename, "scan", None, images=images,
                               note=note or "A scanned PDF: its pages are pictures")
            return DocText(filename, "scan", None, note="A scanned PDF with no readable pages")
        if ext in (".xlsx", ".xlsm") or "spreadsheetml" in ct:
            text, note = _excel(data)
            return DocText(filename, "excel", _clip(text), note=note)
        if ext == ".xls" or ct == "application/vnd.ms-excel":
            return DocText(filename, "other", None,
                           note="Old Excel files (.xls) are not read yet; ask for .xlsx")
        if ext == ".csv" or ct == "text/csv":
            return DocText(filename, "csv", _clip(_csv(data)))
        if ext == ".docx" or "wordprocessingml" in ct:
            return DocText(filename, "word", _clip(_word(data)))
        if ext in (".txt",) or ct == "text/plain":
            return DocText(filename, "text", _clip(_decode(data)))
        if ct in IMAGE_TYPES or ext in (".png", ".jpg", ".jpeg", ".webp", ".gif"):
            if len(data) > MAX_IMAGE_BYTES:
                return DocText(filename, "image", None, note="The picture is too large to read")
            mime = ct if ct in IMAGE_TYPES else (
                "image/png" if ext == ".png" else "image/jpeg")
            return DocText(filename, "image", None, images=[(mime, data)])
        if ext in CAD:
            return DocText(filename, "cad", None, note="A CAD drawing: kept, not read")
    except Exception as exc:  # noqa: BLE001 - a bad file is "not read", never a failure
        log.info("could not read attachment %s: %s", filename, type(exc).__name__)
        return DocText(filename, "other", None, note="The file could not be read")
    return DocText(filename, "other", None, note="This kind of file is kept, not read")


async def read(filename: str, content_type: str, data: bytes) -> DocText:
    try:
        return await asyncio.wait_for(asyncio.to_thread(read_sync, filename, content_type,
                                                        data), READ_SECONDS)
    except TimeoutError:
        return DocText(filename, "other", None, note="Reading the file took too long")
